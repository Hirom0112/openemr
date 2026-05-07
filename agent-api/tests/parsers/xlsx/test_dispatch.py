"""Tests for parsers.xlsx.parse_xlsx end-to-end (Phase 9 Slice 9.5)."""

from __future__ import annotations

import io
import zipfile
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from extractors.schemas import IntakeForm, LabReport, PendingTask
from parsers.xlsx import (
    XlsxMacroRejected,
    XlsxMalformedError,
    XlsxMergedCellsRejected,
    parse_xlsx,
)

pytestmark = [pytest.mark.hard_failure, pytest.mark.clinical_accuracy]

FIXTURE_DIR = Path(__file__).parents[2] / "fixtures" / "w2" / "multimodal" / "xlsx"


def _read(name: str) -> bytes:
    return (FIXTURE_DIR / name).read_bytes()


# --------------------------------------------------------------------------- #
# Full round-trip on p01-chen
# --------------------------------------------------------------------------- #


def test_p01_chen_workbook_full_round_trip():
    raw = _read("p01-chen-workbook.xlsx")
    pw = parse_xlsx(raw, document_reference_id="DOC-p01", patient_id="patient-1")

    # IntakeForm
    assert isinstance(pw.intake_form, IntakeForm)
    assert pw.intake_form.classifier_confidence == 1.0
    assert pw.intake_form.ocr_confidence_range == (1.0, 1.0)
    demo = pw.intake_form.demographics
    assert demo is not None
    assert demo.name.value == "Margaret Chen"
    assert demo.dob.value == "1968-03-12"
    assert demo.sex.value == "F"
    assert demo.mrn.value == "BHS-2847163"
    assert demo.address.value.startswith("2418 Channing Way")

    # NKDA collapses to empty list
    assert pw.intake_form.allergies == []

    # Medications: 4 rows, generic (brand) name shape, dose composed
    meds = pw.intake_form.current_medications
    assert len(meds) == 4
    assert meds[0].name == "atorvastatin (Lipitor)"
    assert meds[0].dose == "40 mg | PO | 1 tab PO daily"
    assert meds[0].indication is not None and meds[0].indication.value == "Hyperlipidemia"
    assert meds[0].prescriber is not None and meds[0].prescriber.value == "Dr. Helen Park, MD"
    assert meds[0].last_filled == date(2026, 4, 10)
    assert meds[0].refills_remaining == 3
    # Anchor citation locator
    assert meds[0].citations[0].field_or_chunk_id == "sheet=Medications|row=2|col=Generic"

    # Labs_Trend wide→long: 4 LabReports, 4 LabValues each = 16 LabValues
    assert len(pw.lab_reports) == 4
    assert all(len(r.values) == 4 for r in pw.lab_reports)
    assert sum(len(r.values) for r in pw.lab_reports) == 16

    # Verify per-date alignment for one analyte (LDL)
    ldls = []
    for r in pw.lab_reports:
        for v in r.values:
            if "ldl" in v.normalized_test_name:
                ldls.append((v.collection_date, v.value, v.abnormal_flag))
    assert ldls == [
        (date(2024, 10, 18), "130", "high"),
        (date(2025, 4, 14), "139", "high"),
        (date(2025, 10, 20), "140", "high"),
        (date(2026, 4, 12), "142", "high"),
    ]

    # All locators carry the date column header
    locators = {v.citations[0].field_or_chunk_id for r in pw.lab_reports for v in r.values}
    assert "sheet=Labs_Trend|row=3|col=2026-04-12" in locators

    # Care_Gaps → PendingTasks
    tasks = pw.pending_tasks
    assert len(tasks) == 7
    overdue = [t for t in tasks if t.status == "OVERDUE"]
    assert len(overdue) == 2
    eye_exam = next(t for t in overdue if "Diabetic eye exam" in t.measure.value)
    assert eye_exam.due_date == date(2025, 10, 30)
    assert eye_exam.last_done == date(2024, 10, 30)

    # Sentinel parse — no warnings on a fully valid fixture
    assert pw.warnings == []


# --------------------------------------------------------------------------- #
# Care_Gaps OVERDUE smoke
# --------------------------------------------------------------------------- #


def test_care_gaps_overdue_status_round_trip():
    raw = _read("p04-kowalski-workbook.xlsx")
    pw = parse_xlsx(raw, document_reference_id="DOC-p04", patient_id="patient-4")
    overdue = [t for t in pw.pending_tasks if t.status == "OVERDUE"]
    assert any(t.measure.value.startswith("Hepatitis C") for t in overdue)
    # The HCV row has "—" for last_done — should soft-coerce to None
    hcv = next(t for t in overdue if t.measure.value.startswith("Hepatitis C"))
    assert hcv.last_done is None
    assert hcv.due_date == date(2026, 5, 10)


# --------------------------------------------------------------------------- #
# Labs_Trend wide→long count check across multiple fixtures
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize(
    "fixture_name,expected_dates",
    [
        ("p01-chen-workbook.xlsx", 4),
        ("p02-whitaker-workbook.xlsx", 3),
        ("p04-kowalski-workbook.xlsx", 3),
    ],
)
def test_labs_trend_one_report_per_date_column(fixture_name, expected_dates):
    raw = _read(fixture_name)
    pw = parse_xlsx(raw, document_reference_id="DOC-x", patient_id="patient-x")
    assert len(pw.lab_reports) == expected_dates
    # Each LabReport has at least one value
    assert all(len(r.values) >= 1 for r in pw.lab_reports)
    # Citations always carry a Labs_Trend locator with a date col
    for r in pw.lab_reports:
        for v in r.values:
            assert v.citations[0].field_or_chunk_id.startswith("sheet=Labs_Trend|row=")
            assert "|col=" in v.citations[0].field_or_chunk_id


# --------------------------------------------------------------------------- #
# Macro / merged-cell rejection
# --------------------------------------------------------------------------- #


def _inject_vba(raw: bytes) -> bytes:
    """Return a copy of ``raw`` with a stub xl/vbaProject.bin entry added."""
    bio = io.BytesIO(raw)
    out = io.BytesIO()
    with zipfile.ZipFile(bio, "r") as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.namelist():
            dst.writestr(item, src.read(item))
        dst.writestr("xl/vbaProject.bin", b"\x00\x00stub")
    return out.getvalue()


def test_macro_xlsm_rejected_at_dispatch():
    raw = _read("p01-chen-workbook.xlsx")
    macro_raw = _inject_vba(raw)
    with pytest.raises(XlsxMacroRejected):
        parse_xlsx(macro_raw, document_reference_id="DOC-x", patient_id="patient-x")


def test_merged_cells_rejected():
    """Merge two cells in the Patient sheet then re-pack — expect rejection."""
    raw = _read("p01-chen-workbook.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    ws = wb["Patient"]
    ws.merge_cells("A1:B1")
    bio = io.BytesIO()
    wb.save(bio)
    with pytest.raises(XlsxMergedCellsRejected):
        parse_xlsx(bio.getvalue(), document_reference_id="DOC-x", patient_id="patient-x")


# --------------------------------------------------------------------------- #
# Missing-sheet soft-warn / hard fail
# --------------------------------------------------------------------------- #


def _strip_sheets(raw: bytes, keep: set[str]) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    for name in list(wb.sheetnames):
        if name not in keep:
            del wb[name]
    if not wb.sheetnames:
        # openpyxl requires at least one sheet
        wb.create_sheet("Empty")
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def test_missing_one_sheet_soft_warns():
    raw = _read("p01-chen-workbook.xlsx")
    stripped = _strip_sheets(raw, keep={"Patient", "Medications", "Labs_Trend"})
    pw = parse_xlsx(stripped, document_reference_id="DOC-x", patient_id="patient-x")
    assert any(w == "xlsx_sheet_missing:Care_Gaps" for w in pw.warnings)
    assert pw.intake_form is not None
    assert pw.lab_reports  # still extract labs
    assert pw.pending_tasks == []


def test_all_known_sheets_missing_hard_fails():
    raw = _read("p01-chen-workbook.xlsx")
    stripped = _strip_sheets(raw, keep=set())
    with pytest.raises(XlsxMalformedError) as excinfo:
        parse_xlsx(stripped, document_reference_id="DOC-x", patient_id="patient-x")
    assert excinfo.value.code == "xlsx_no_known_sheets"


# --------------------------------------------------------------------------- #
# Alias-table drift
# --------------------------------------------------------------------------- #


def test_patient_sheet_address_aliases():
    """Mailing_Address / Home_Address collapse to Address."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Patient"
    ws["A1"] = "Field"
    ws["B1"] = "Value"
    ws["A2"] = "Name"
    ws["B2"] = "Test Patient"
    ws["A3"] = "MRN"
    ws["B3"] = "MRN-1"
    ws["A4"] = "DOB"
    ws["B4"] = "1990-01-01"
    ws["A5"] = "Mailing_Address"
    ws["B5"] = "1 Main St"
    bio = io.BytesIO()
    wb.save(bio)
    pw = parse_xlsx(bio.getvalue(), document_reference_id="DOC-x", patient_id="patient-x")
    assert pw.intake_form is not None
    assert pw.intake_form.demographics.address is not None
    assert pw.intake_form.demographics.address.value == "1 Main St"


# --------------------------------------------------------------------------- #
# Date-as-serial-number handling for Medications.last_filled
# --------------------------------------------------------------------------- #


def test_medications_last_filled_handles_datetime_object():
    """openpyxl turns date-formatted cells into datetime; we coerce to date."""
    from datetime import datetime as _dt
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Medications"
    ws["A1"] = "Brand"
    ws["B1"] = "Generic"
    ws["C1"] = "Strength"
    ws["D1"] = "Last_Filled"
    ws["E1"] = "Refills_Remaining"
    ws["A2"] = "Lipitor"
    ws["B2"] = "atorvastatin"
    ws["C2"] = "20 mg"
    ws["D2"] = _dt(2025, 7, 15, 12, 0, 0)
    ws["E2"] = 4
    # Add stub Patient sheet to satisfy known-sheet gate
    ws2 = wb.create_sheet("Patient")
    ws2["A1"] = "Field"
    ws2["B1"] = "Value"
    ws2["A2"] = "Name"
    ws2["B2"] = "Test"
    bio = io.BytesIO()
    wb.save(bio)
    pw = parse_xlsx(bio.getvalue(), document_reference_id="DOC-x", patient_id="patient-x")
    meds = pw.intake_form.current_medications
    assert len(meds) == 1
    assert meds[0].last_filled == date(2025, 7, 15)
    assert meds[0].refills_remaining == 4


# --------------------------------------------------------------------------- #
# Type checks
# --------------------------------------------------------------------------- #


def test_returns_typed_lab_reports_and_tasks():
    raw = _read("p01-chen-workbook.xlsx")
    pw = parse_xlsx(raw, document_reference_id="DOC-p01", patient_id="patient-1")
    assert all(isinstance(r, LabReport) for r in pw.lab_reports)
    assert all(isinstance(t, PendingTask) for t in pw.pending_tasks)
