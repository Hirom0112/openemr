"""Tests for parsers.xlsx.probe (Phase 9 Slice 9.5)."""

from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pytest

from parsers.xlsx import probe_identity
from parsers.xlsx.types import CandidateHints

pytestmark = [pytest.mark.hard_failure, pytest.mark.clinical_accuracy]

FIXTURE_DIR = Path(__file__).parents[2] / "fixtures" / "w2" / "multimodal" / "xlsx"


def test_probe_identity_extracts_p01_chen():
    raw = (FIXTURE_DIR / "p01-chen-workbook.xlsx").read_bytes()
    hints = probe_identity(raw)
    assert isinstance(hints, CandidateHints)
    assert hints.mrn == "BHS-2847163"
    assert hints.name_family == "Chen"
    assert hints.name_given == "Margaret"
    assert hints.dob == "1968-03-12"


def test_probe_identity_returns_none_on_non_xlsx_bytes():
    assert probe_identity(b"not an xlsx") is None


def test_probe_identity_returns_none_when_no_patient_sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Other"
    ws["A1"] = "x"
    bio = io.BytesIO()
    wb.save(bio)
    assert probe_identity(bio.getvalue()) is None


def test_probe_identity_handles_family_comma_given():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Patient"
    ws["A1"] = "Field"
    ws["B1"] = "Value"
    ws["A2"] = "Name"
    ws["B2"] = "Smith, John"
    ws["A3"] = "MRN"
    ws["B3"] = "MRN-7"
    ws["A4"] = "DOB"
    ws["B4"] = "1980-04-22"
    bio = io.BytesIO()
    wb.save(bio)
    hints = probe_identity(bio.getvalue())
    assert hints is not None
    assert hints.name_family == "Smith"
    assert hints.name_given == "John"
    assert hints.mrn == "MRN-7"
    assert hints.dob == "1980-04-22"
